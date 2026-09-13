#!/usr/bin/env python3
"""Validate the canonical Archie user-story workbook before review or retest."""

from pathlib import Path
import re
import sys

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
WORKBOOK = ROOT / "docs/verification/user-stories.xlsx"
STATUSES = {"Not tested", "Partially verified", "Pass", "Fail", "Blocked", "Intentionally unavailable"}
RETEST_STATUSES = STATUSES | {"Not run"}
STORY_HEADERS = {
    "ID", "Area", "Feature", "User Story", "Expected Behavior", "Source Refs",
    "Prerequisites", "Test Steps", "Verification Method", "Verification Level",
    "Current Status", "Errors", "Error Severity", "Fix", "Retest Status", "Evidence",
    "Code Review Status", "Baseline Status", "Baseline Evidence", "Post-Fix Status",
    "Post-Fix Evidence",
}


def fail(message: str) -> None:
    print(f"ERROR: {message}", file=sys.stderr)
    raise SystemExit(1)


wb = load_workbook(WORKBOOK, read_only=True, data_only=True)
if "Stories" not in wb or "Coverage" not in wb or "Errors" not in wb:
    fail("workbook must contain Stories, Coverage, and Errors sheets")

stories = wb["Stories"]
headers = [cell.value for cell in next(stories.iter_rows(values_only=False))]
missing = STORY_HEADERS - set(headers)
if missing:
    fail(f"Stories is missing columns: {sorted(missing)}")
col = {name: headers.index(name) for name in headers}
story_ids: list[str] = []
story_statuses: dict[str, str] = {}
for row_number, row in enumerate(stories.iter_rows(min_row=2, values_only=True), 2):
    if not any(value is not None for value in row):
        continue
    story_id = row[col["ID"]]
    if not story_id:
        fail(f"Stories row {row_number} has no ID")
    if story_id in story_ids:
        fail(f"duplicate story ID: {story_id}")
    story_ids.append(story_id)
    story_statuses[story_id] = row[col["Current Status"]]
    for field in ("Area", "Feature", "User Story", "Expected Behavior", "Source Refs", "Prerequisites", "Test Steps", "Verification Method"):
        if not row[col[field]]:
            fail(f"{story_id} has blank {field}")
    if row[col["Current Status"]] not in STATUSES:
        fail(f"{story_id} has invalid current status {row[col['Current Status']]!r}")
    for status_field in ("Baseline Status", "Post-Fix Status", "Retest Status"):
        status_value = row[col[status_field]]
        if status_value and status_value not in RETEST_STATUSES:
            fail(f"{story_id} has invalid {status_field} {status_value!r}")
    for status_field, evidence_field in (("Current Status", "Evidence"), ("Baseline Status", "Baseline Evidence"), ("Post-Fix Status", "Post-Fix Evidence")):
        status_value = row[col[status_field]]
        if status_value == "Pass" and not row[col[evidence_field]]:
            fail(f"{story_id} is Pass in {status_field} but has no {evidence_field}")
    prerequisites = str(row[col["Prerequisites"]])
    if re.search(r"(?:^|\n)[A-Za-z0-9 .,;:/()'’!?-](?:\n[A-Za-z0-9 .,;:/()'’!?-]){4,}(?:$|\n)", prerequisites):
        fail(f"{story_id} Prerequisites appears character-per-line corrupted")
    for ref in str(row[col["Source Refs"]]).split("\n"):
        match = re.match(r"^(.*?):(\d+)(?:-(\d+))?$", ref)
        if not match:
            fail(f"{story_id} has malformed source ref {ref!r}")
        source = ROOT / match.group(1)
        if not source.is_file():
            fail(f"{story_id} references missing source {match.group(1)}")
        lines = len(source.read_text(errors="replace").splitlines())
        start = int(match.group(2))
        end = int(match.group(3) or start)
        if not 1 <= start <= end <= lines:
            fail(f"{story_id} source range exceeds {source}: {start}-{end} of {lines}")

covered: list[str] = []
coverage = wb["Coverage"]
for row in coverage.iter_rows(min_row=2, values_only=True):
    if row[2]:
        covered.extend(item.strip() for item in str(row[2]).split(",") if item.strip())
if set(covered) != set(story_ids) or len(covered) != len(set(covered)):
    fail("Coverage must map every story exactly once")

errors = wb["Errors"]
error_impacted: set[str] = set()
for row_number, row in enumerate(errors.iter_rows(min_row=2, values_only=True), 2):
    if not any(value is not None for value in row):
        continue
    impacted = [item.strip() for item in str(row[1] or "").split(",") if item.strip()]
    unknown = set(impacted) - set(story_ids)
    if unknown:
        fail(f"Errors row {row_number} references unknown stories: {sorted(unknown)}")
    error_impacted.update(impacted)
    if not row[10]:
        fail(f"Errors row {row_number} has no Category")

for story_id, status in story_statuses.items():
    if status == "Fail" and story_id not in error_impacted:
        fail(f"{story_id} is Fail but has no linked Errors row")

print(f"Validated {len(story_ids)} stories, {len(covered)} coverage mappings, and {errors.max_row - 1} error rows.")
